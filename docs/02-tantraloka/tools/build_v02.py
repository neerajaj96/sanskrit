#!/usr/bin/env python3
"""Tantraloka Malayalam Series — v02 production builder (A4 + Trade).

Implements SERIES_STYLE.md: original house style, true footnotes, curated
figures, per-chapter running heads, roman/arabic folios, two-pass TOC with
real page numbers, locked 3-family font architecture, PDF bookmarks via
heading styles + explicit anchors.

Usage:
  python3 build_v02.py --edition a4            # pass 1 (TOC without numbers)
  python3 build_v02.py --edition a4 --toc TOC.json   # pass 2 (final)
  python3 build_v02.py --edition trade [--toc ...]
Outputs (never touches v01):
  pdf/Tantraloka_Malayalam_Volume1_v02_A4.docx / _A4.pdf
  pdf/Tantraloka_Malayalam_Volume1_v02_Trade.docx / _Trade.pdf
"""
import argparse
import copy
import json
import os
import re

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = "/root/tantraloka_ml"
TR = os.path.join(BASE, "translated")
FIGDIR = os.path.join(BASE, "assets", "figs_v02")
EMBLEM = os.path.join(BASE, "assets", "cover_emblem.png")
CAPJSON = os.path.join(BASE, "assets", "captions.json")
try:
    with open(CAPJSON, encoding="utf-8") as _cf:
        FIGCAPS = json.load(_cf)
except Exception:
    FIGCAPS = {}

# ---------------- Series design tokens (see SERIES_STYLE.md) ----------------
IVORY = RGBColor(0xFF, 0xFB, 0xF0)
GOLD = RGBColor(0xA6, 0x7C, 0x2E)
GOLD_LT = RGBColor(0xC4, 0x9E, 0x50)
MAROON = RGBColor(0x5A, 0x1A, 0x10)
MAROON_DK = RGBColor(0x4A, 0x14, 0x0C)
SLATE = RGBColor(0x2E, 0x40, 0x57)
INK = RGBColor(0x21, 0x21, 0x21)
GREY = RGBColor(0x6E, 0x6E, 0x6E)

BODY_FONT = "Noto Serif Malayalam"
HEAD_FONT = "Noto Sans Malayalam"
LATIN_FONT = "Noto Serif"          # IAST / Latin: full diacritic coverage

_MLD = "൦൧൨൩൪൫൬൭൮൯"

LATIN_RUN = re.compile(r"[A-Za-z\u0100-\u024F\u1E00-\u1EFF\u2C60-\u2C7F\u00B7\u2022\u2026]+")
# NB: U+00B7 MIDDLE DOT and U+2022 BULLET are absent from the Noto Malayalam
# fonts (verified via fontTools); routing them to Noto Serif kills fallback.


def m2i(s):
    v = 0
    for ch in s:
        if ch in _MLD:
            v = v * 10 + _MLD.index(ch)
    return v


EDITIONS = {
    "a4": dict(
        page_w=Inches(8.27), page_h=Inches(11.69),
        m_top=Inches(0.75), m_bot=Inches(0.80),
        m_left=Inches(0.90), m_right=Inches(0.90),
        body_pt=10.8, mula_pt=11.5, note_pt=8.8, cap_pt=8.5,
        fig_w=Inches(2.9), emblem_w=Inches(3.0),
        line_mult=1.42, head_pt=7.5, series_head="തന്ത്രാലോകം — മലയാളപതിപ്പ്",
        suffix="A4",
    ),
    "trade": dict(
        page_w=Inches(6.0), page_h=Inches(9.0),
        m_top=Inches(0.65), m_bot=Inches(0.70),
        m_left=Inches(0.65), m_right=Inches(0.65),
        body_pt=10.2, mula_pt=10.8, note_pt=8.4, cap_pt=8.0,
        fig_w=Inches(2.5), emblem_w=Inches(2.4),
        line_mult=1.40, head_pt=7.0, series_head="തന്ത്രാലോകം",
        suffix="Trade",
    ),
}

ORDER = [
    ("ta_v02.md", "ശ്ലോകങ്ങൾ ൧–൨൧ — ആമുഖവന്ദനം",
     "ആമുഖവന്ദനം · ൧–൨൧", "V001"),
    ("ta_v03.md", "ശ്ലോകങ്ങൾ ൨൨–൧൦൫ — ആമുഖം",
     "ആമുഖം · ൨൨–൧൦൫", "V022"),
    ("ta_v04.md", "ശ്ലോകങ്ങൾ ൧൦൬–൧൩൯ — വിവരണം",
     "വിവരണം · ൧൦൬–൧൩൯", "V106"),
    ("ta_v05.md", "ശ്ലോകങ്ങൾ ൧൪൦–൨൪൫ — സാധനമാർഗങ്ങൾ",
     "സാധനമാർഗ്ഗങ്ങൾ · ൧൪൦–൨൪൫", "V140"),
    ("ta_v06.md", "ശ്ലോകങ്ങൾ ൨൪൭–൩൩൫ — വിഷയനിർദ്ദേശം",
     "വിഷയനിർദ്ദേശം · ൨൪൭–൩൩൫", "V247"),
    ("ta_app.md", "അനുബന്ധങ്ങൾ — ശബ്ദതലങ്ങൾ · ദ്വാദശാന്തം",
     "അനുബന്ധങ്ങൾ", "APP"),
]

CURATED_VERSES = {1, 2, 7, 14, 22, 36, 39, 52, 60, 67, 81, 90, 94, 106,
                  116, 125, 134, 140, 150, 156, 161, 167, 171, 202, 211,
                  214, 219, 226, 232, 239, 247, 274, 279, 288, 331}


# ---------------- low-level helpers ----------------
def set_run_font(run, family, size=None, bold=None, italic=None, color=None):
    run.font.name = family
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), family)


def add_mixed_text(paragraph, text, ml_font, base_size=None, base_bold=False,
                   base_italic=False, color=None, latin_bold=None,
                   latin_italic=None):
    """Segment-wise font locking: Latin/IAST runs -> LATIN_FONT, all else
    Malayalam family. Kills arbitrary fallback for diacritics."""
    pos = 0
    for m in LATIN_RUN.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_run_font(r, ml_font, base_size, base_bold, base_italic, color)
        r = paragraph.add_run(m.group(0))
        set_run_font(r, LATIN_FONT, base_size,
                     base_bold if latin_bold is None else latin_bold,
                     base_italic if latin_italic is None else latin_italic,
                     color)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_run_font(r, ml_font, base_size, base_bold, base_italic, color)


def add_rich_mixed(paragraph, text, ml_font, base_size=None, color=None):
    """**bold** / *italic* markup + script-aware fonts."""
    tokens = re.split(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            add_mixed_text(paragraph, tok[2:-2], ml_font, base_size,
                           True, False, color)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            add_mixed_text(paragraph, tok[1:-1], ml_font, base_size,
                           False, True, color)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            add_mixed_text(paragraph, tok[1:-1], ml_font, base_size,
                           False, False, color)
        else:
            add_mixed_text(paragraph, tok, ml_font, base_size,
                           False, False, color)


def shade(paragraph, hex_fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    pPr.append(shd)


def box(paragraph, color_hex="A67C2E", size="4"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:color'), color_hex)
        pBdr.append(el)
    pPr.append(pBdr)


def hline(paragraph, color="A67C2E", size="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def keep_lines(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:keepLines'))


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:keepNext'))

def styled_heading(doc, text, level, anchor=None, toc=None, toc_level=None,
                   size=None, color=None, align=None, latin_italic_heading=False):
    """Heading with explicit per-run fonts (LibreOffice-safe CTL)."""
    h = doc.add_heading(level=level)
    h.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    pos = 0
    for m in LATIN_RUN.finditer(text):
        if m.start() > pos:
            r = h.add_run(text[pos:m.start()])
            set_run_font(r, HEAD_FONT, size, True, False, color)
        r = h.add_run(m.group(0))
        set_run_font(r, LATIN_FONT, size, True, latin_italic_heading, color)
        pos = m.end()
    if pos < len(text):
        r = h.add_run(text[pos:])
        set_run_font(r, HEAD_FONT, size, True, False, color)
    if anchor:
        add_bookmark(h, anchor)
    if toc is not None and anchor:
        toc.append((toc_level or level, text, anchor))
    return h



_bookmark_id = [2000]


class AnchorGen:
    """Deterministic unique anchor factory (shared by scanner + renderer)."""

    def __init__(self):
        self.used = set()

    def slug(self, t):
        return re.sub(r'\s+', '_', re.sub(
            r'[^\w൦-൯]+', ' ', t, flags=re.UNICODE)).strip('_')[:20]

    def get(self, base):
        if base not in self.used:
            self.used.add(base)
            return base
        k = 2
        while f"{base}_{k}" in self.used:
            k += 1
        self.used.add(f"{base}_{k}")
        return f"{base}_{k}"

    def h2(self, code, t):
        m = re.search(r'[൦-൯]+', t)
        base = (f"ch_{code}_V{m.group(0)}" if m
                else f"ch_{code}_{self.slug(t)}")
        return self.get(base)

    def recap(self, code, t):
        m = re.search(r'[൦-൯]+', t)
        base = (f"ch_{code}_R{m.group(0)}" if m
                else f"ch_{code}_R{self.slug(t)}")
        return self.get(base)


def scan_headings(AG=None):
    """Pure pre-scan of content order/titles/anchors for the TOC (spec M:
    TOC built from final structure before frontmatter is rendered)."""
    AG = AG or AnchorGen()
    out = []
    allanchors = []
    for fname, title, short, code in ORDER:
        if code == "APP":
            continue
        path = os.path.join(TR, fname)
        with open(path, encoding="utf-8") as f:
            md = f.read().splitlines()
        _a = AG.get(f"ch_{code}")
        allanchors.append(_a)
        out.append((1, title, _a))
        for line in md:
            s = line.strip()
            if s.startswith('# ') and not s.startswith('## '):
                # file-internal H1: skipped on-page (VQ-003); no anchor.
                continue
            elif s.startswith('## '):
                t = s[3:].strip()
                if t.startswith('ഇതിൽനിന്ന്'):
                    allanchors.append(AG.recap(code, t))  # not in TOC
                else:
                    _a = AG.h2(code, t)
                    allanchors.append(_a)
                    out.append((2, t, _a))
    # appendices (mirror render_appendices)
    path = os.path.join(TR, "ta_app.md")
    with open(path, encoding="utf-8") as f:
        md = f.read().splitlines()
    _a = AG.get("ch_APP")
    allanchors.append(_a)
    out.append((1, "അനുബന്ധങ്ങൾ — ശബ്ദതലങ്ങൾ · ദ്വാദശാന്തം", _a))
    for line in md:
        s = line.strip()
        if s.startswith('## '):
            t = s[3:].strip()
            if t.startswith('അനുബന്ധം'):
                parts = t.split('—', 1)
                code2 = ('APP1' if '൧' in parts[0]
                         else ('APP2' if '൨' in parts[0] else 'APPG'))
                _a = AG.get(f"ch_{code2}")
                allanchors.append(_a)
                out.append((1, t, _a))
            else:
                _a = AG.get("ch_APPG")
                allanchors.append(_a)
                out.append((2, t, _a))
    return out, allanchors


def add_bookmark(paragraph, name):
    _bookmark_id[0] += 1
    bid = str(_bookmark_id[0])
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), bid)
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), bid)
    p = paragraph._p
    p.insert(0, start)
    p.append(end)


def add_internal_link(paragraph, bookmark, text, size=None):
    hl = OxmlElement('w:hyperlink')
    hl.set(qn('w:anchor'), bookmark)
    hl.set(qn('w:history'), '1')
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rF = OxmlElement('w:rFonts')
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rF.set(qn(attr), HEAD_FONT)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '2E4057')
    if size is not None:
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(size.pt * 2)))
        rPr.append(sz)
    rPr.append(rF)
    rPr.append(color)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    hl.append(run)
    paragraph._p.append(hl)


def add_toc_dot_leader(paragraph):
    """Right tab with dot leader for TOC page numbers."""
    pPr = paragraph._p.get_or_add_pPr()
    tabs = pPr.find(qn('w:tabs'))
    if tabs is None:
        tabs = OxmlElement('w:tabs')
        pPr.append(tabs)
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), '9350')
    tabs.append(tab)


# ---------------- footnotes ----------------
class Footnotes:
    FN_RT = ('http://schemas.openxmlformats.org/officeDocument/'
             '2006/relationships/footnotes')

    def __init__(self, doc, note_size):
        self.doc = doc
        self.note_size = note_size
        self.counter = [1]
        self.footnotes_el = OxmlElement('w:footnotes')
        for fid, typ in (('0', 'separator'), ('1', 'continuationSeparator')):
            fn = OxmlElement('w:footnote')
            fn.set(qn('w:id'), fid)
            fn.set(qn('w:type'), typ)
            par = OxmlElement('w:p')
            run = OxmlElement('w:r')
            sep = OxmlElement('w:separator') if typ == 'separator' else \
                OxmlElement('w:continuationSeparator')
            run.append(sep)
            par.append(run)
            fn.append(par)
            self.footnotes_el.append(fn)
        self._installed = False

    def _install(self):
        if self._installed:
            return
        from lxml import etree
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        xml_bytes = etree.tostring(self.footnotes_el, xml_declaration=True,
                                   encoding='UTF-8', standalone=True)
        uri = PackURI('/word/footnotes.xml')
        fpart = Part(uri, 'application/vnd.openxmlformats-officedocument.'
                          'wordprocessingml.footnotes+xml',
                     xml_bytes, self.doc.part.package)
        self.doc.part.relate_to(fpart, self.FN_RT)
        self._part = fpart
        self._installed = True

    def add(self, anchor_paragraph, text):
        """Attach footnote; anchor appended to anchor_paragraph."""
        from lxml import etree
        fid = str(self.counter[0] + 1)
        self.counter[0] += 1
        fn = OxmlElement('w:footnote')
        fn.set(qn('w:id'), fid)
        par = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), 'both')
        pPr.append(jc)
        par.append(pPr)
        # build runs manually with mixed fonts
        # tokenize bold lead
        m = re.match(r'(\*\*.+?\*\*)\s*(.*)', text, re.S)
        segs = []
        if m:
            segs.append((m.group(1)[2:-2], True, False))
            text = m.group(2)
        # split remaining **..** occurrences
        for tok in re.split(r'(\*\*.+?\*\*)', text):
            if not tok:
                continue
            if tok.startswith('**') and tok.endswith('**'):
                segs.append((tok[2:-2], True, False))
            else:
                segs.append((tok, False, False))
        for seg_text, b, it in segs:
            pos = 0
            for mm in LATIN_RUN.finditer(seg_text):
                if mm.start() > pos:
                    self._note_run(par, seg_text[pos:mm.start()],
                                   BODY_FONT, b, it)
                self._note_run(par, mm.group(0), LATIN_FONT, b, it)
                pos = mm.end()
            if pos < len(seg_text):
                self._note_run(par, seg_text[pos:], BODY_FONT, b, it)
        fn.append(par)
        self.footnotes_el.append(fn)
        # reference run on anchor
        r = anchor_paragraph.add_run()
        rPr = r._r.get_or_add_rPr()
        st = OxmlElement('w:rStyle')
        st.set(qn('w:val'), 'FootnoteReference')
        rPr.append(st)
        ref = OxmlElement('w:footnoteReference')
        ref.set(qn('w:id'), fid)
        r._r.append(ref)

    def _note_run(self, par, text, family, bold, italic):
        if not text:
            return
        run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rF = OxmlElement('w:rFonts')
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rF.set(qn(attr), family)
        rPr.append(rF)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(self.note_size.pt * 2)))
        rPr.append(sz)
        if bold:
            rPr.append(OxmlElement('w:b'))
        if italic:
            rPr.append(OxmlElement('w:i'))
        run.append(rPr)
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        run.append(t)
        # footnote text lives in w:p of the footnote
        par.append(run)

    def finalize(self):
        from lxml import etree
        self._install()
        self._part._blob = etree.tostring(
            self.footnotes_el, xml_declaration=True, encoding='UTF-8',
            standalone=True)


# ---------------- tables ----------------
def add_word_table(doc, header_cells, body_rows, body_pt, caption=None,
                   widths=None):
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(2)
        add_mixed_text(cp, caption, HEAD_FONT, Pt(9), True, False, MAROON)
        keep_with_next(cp)
    table = doc.add_table(rows=1 + len(body_rows), cols=len(header_cells))
    table.style = 'Table Grid'
    table.autofit = False
    n = len(header_cells)
    if widths is None:
        if n == 4:
            widths = [1.9, 1.5, 1.7, 1.3]
        elif n == 3:
            widths = [1.1, 2.6, 2.7]
        else:
            widths = [6.4 / n] * n
    s = sum(widths)
    widths = [w / s * 6.4 for w in widths]
    for j, w in enumerate(widths):
        for row in table.rows:
            row.cells[j].width = Inches(w)
    for j, h in enumerate(header_cells):
        cell = table.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        add_mixed_text(p, h, HEAD_FONT, Pt(9), True, False,
                       RGBColor(0xFF, 0xFF, 0xFF))
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '5A1A10')
        tcPr.append(shd)
    # repeat header row on page breaks
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    for i, row in enumerate(body_rows):
        for j in range(len(header_cells)):
            txt = row[j] if j < len(row) else "—"
            cell = table.cell(i + 1, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_mixed_text(p, txt, BODY_FONT, body_pt, False, False, INK)
            if i % 2 == 1:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'FFFBF0')
                tcPr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def parse_md_row(s):
    return [c.strip() for c in s.strip().strip('|').split('|')]


# ---------------- document-wide styles ----------------
def style_document(doc, E):
    style = doc.styles['Normal']
    style.font.name = BODY_FONT
    style.font.size = Pt(E["body_pt"])
    style.font.color.rgb = INK
    pf = style.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing = E["line_mult"]
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True
    pf.first_line_indent = Inches(0.22)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rFonts.set(qn(attr), BODY_FONT)
    specs = {1: (16.5, MAROON_DK, 10, 4), 2: (12.5, MAROON, 8, 3),
             3: (11.0, SLATE, 6, 3)}
    for i, (size, color, before, after) in specs.items():
        hs = doc.styles[f'Heading {i}']
        hs.font.name = HEAD_FONT
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = color
        hs.paragraph_format.keep_with_next = True
        hs.paragraph_format.widow_control = True
        hs.paragraph_format.space_before = Pt(before)
        hs.paragraph_format.space_after = Pt(after)
        hs.paragraph_format.first_line_indent = Inches(0)
        if i > 1:
            hs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rPr2 = hs.element.get_or_add_rPr()
        rf = rPr2.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rPr2.append(rf)
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rf.set(qn(attr), HEAD_FONT)


def setup_section(sec, E, header_odd=None, header_even=None,
                  footer_fmt="decimal", footer_start=None,
                  first_empty=True, page_border=False, folio=True,
                  restart_footnotes=False):
    sec.page_height = E["page_h"]
    sec.page_width = E["page_w"]
    sec.top_margin = E["m_top"]
    sec.bottom_margin = E["m_bot"]
    sec.left_margin = E["m_left"]
    sec.right_margin = E["m_right"]
    sec.odd_and_even_pages_header_footer = True
    sec.different_first_page_header_footer = first_empty
    # Each section needs INDEPENDENT header/footer parts; otherwise Word
    # links to previous and our runs accumulate across chapters (spec N/O).
    for _part in (sec.header, sec.footer):
        try:
            _part.is_linked_to_previous = False
        except Exception:
            pass
    if header_even is not None:
        try:
            sec.even_page_header.is_linked_to_previous = False
        except Exception:
            pass
    try:
        sec.first_page_header.is_linked_to_previous = False
    except Exception:
        pass
    try:
        sec.first_page_footer.is_linked_to_previous = False
    except Exception:
        pass
    # NOTE: no w:pgBorders page frames anywhere (spec O — restrained folios).
    # LibreOffice propagates the first section's page border to all pages,
    # which overdecorates the book; the cover is framed by its emblem plate.
    if footer_start is not None:
        sectPr = sec._sectPr
        pgNum = OxmlElement('w:pgNumType')
        pgNum.set(qn('w:start'), str(footer_start))
        pgNum.set(qn('w:fmt'), footer_fmt)
        sectPr.append(pgNum)
    # NOTE: w:footnotePr numRestart=eachSect makes LibreOffice render
    # every marker after the first as "0" (verified visually); global
    # numbering renders correctly, so restart is intentionally NOT set.
    # Footnote TITLE labels (കുറിപ്പ് ൧..) preserve the source per-part
    # indexing; markers are document-unique anchors. See BUILD_LOG.
    if False and restart_footnotes:
        pass
    # Single centered running head carrying BOTH series + section identity.
    # (Word even/odd heads are ignored by the LibreOffice converter — verified
    # by spike — so parity-split Left/Right heads are not achievable in this
    # toolchain; combined head preserves the spec-N information. See BUILD_LOG.)
    if header_odd:
        hp = sec.header.paragraphs[0]
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.first_line_indent = Inches(0)
        add_mixed_text(hp, header_odd, HEAD_FONT, Pt(E.get("head_pt", 7.5)),
                       False, False, GOLD)
    if header_even:
        hp = sec.even_page_header.paragraphs[0]
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.first_line_indent = Inches(0)
        add_mixed_text(hp, header_even, HEAD_FONT, Pt(7.5), False, False,
                       GOLD)
    # footer folio
    if not folio:
        try:
            sec.footer.paragraphs[0].clear()
            sec.header.paragraphs[0].clear()
        except Exception:
            pass
        return
    fp = sec.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Inches(0)
    run = fp.add_run()
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1)
    run2 = fp.add_run()
    run2._r.append(instr)
    run3 = fp.add_run()
    run3._r.append(fld2)
    for r in fp.runs:
        set_run_font(r, HEAD_FONT, Pt(8), False, False, GREY)
    if first_empty:
        # thin folio on first pages too (reduced): leave default (empty
        # first header, but first footer keeps folio via primary footer?
        # Word uses separate first-page footer; copy folio there.)
        try:
            ffp = sec.first_page_footer.paragraphs[0]
            ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = ffp.add_run()
            ff1 = OxmlElement('w:fldChar')
            ff1.set(qn('w:fldCharType'), 'begin')
            ii = OxmlElement('w:instrText')
            ii.set(qn('xml:space'), 'preserve')
            ii.text = 'PAGE'
            ff2 = OxmlElement('w:fldChar')
            ff2.set(qn('w:fldCharType'), 'end')
            rr._r.append(ff1)
            rr2 = ffp.add_run()
            rr2._r.append(ii)
            rr3 = ffp.add_run()
            rr3._r.append(ff2)
            for r in ffp.runs:
                set_run_font(r, HEAD_FONT, Pt(8), False, False, GREY)
        except Exception:
            pass


# ---------------- content rendering ----------------
def render_mula_group(doc, E, lines):
    """One compact box per verse group (multi-line verses stay together)."""
    text = "\n".join(l.lstrip('> ').strip() for l in lines if l.strip('> '))
    text = re.sub(r'\n{2,}', '\n', text).strip()
    if not text:
        return None
    pp = doc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.space_before = Pt(4)
    pp.paragraph_format.space_after = Pt(6)
    pp.paragraph_format.first_line_indent = Inches(0)
    shade(pp, "FFFBF0")
    box(pp, "A67C2E", "4")
    keep_lines(pp)
    parts = text.split('\n')
    for k, part in enumerate(parts):
        r = pp.add_run(part)
        set_run_font(r, BODY_FONT, Pt(E["mula_pt"]), True, False, MAROON)
        if k < len(parts) - 1:
            r.add_break()
    return pp


def render_figure_unit(doc, E, verse, caption_text=None):
    """Figure + caption as one controlled visual unit (spec G)."""
    png = os.path.join(FIGDIR, "fig_V%03d.png" % verse)
    if not os.path.exists(png):
        return
    key = str(verse)
    cap = caption_text or FIGCAPS.get(
        key, "ചിത്രം %s" % "".join(_MLD[int(d)] for d in str(verse)))
    pp = doc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.space_before = Pt(4)
    pp.paragraph_format.space_after = Pt(1)
    pp.paragraph_format.first_line_indent = Inches(0)
    keep_lines(pp)
    try:
        pp.add_run().add_picture(png, width=E["fig_w"])
    except Exception:
        return
    keep_with_next(pp)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(1)
    cp.paragraph_format.space_after = Pt(6)
    cp.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(cp, cap, HEAD_FONT, Pt(E["cap_pt"]), False, True, MAROON)
    keep_lines(cp)


def body_para(doc, E, text, footnote_text=None, F=None, no_indent=False,
              size=None, italic=False):
    pp = doc.add_paragraph()
    if no_indent:
        pp.paragraph_format.first_line_indent = Inches(0)
    add_rich_mixed(pp, text, BODY_FONT, size, INK)
    if italic:
        for r in pp.runs:
            r.italic = True
    if footnote_text and F is not None:
        F.add(pp, footnote_text)
    return pp


LEAD_LABELS = ("**നേരർത്ഥം:**", "**ആചാര്യവ്യാഖ്യാനം:**", "**വിശകലനം:**",
               "**കുറിപ്പ്")


def render_chapter_file(doc, E, fname, title, short, code, F, AG):
    path = os.path.join(TR, fname)
    with open(path, encoding="utf-8") as f:
        md = f.read().splitlines()
    styled_heading(doc, title, 1, anchor=AG.get(f"ch_{code}"),
                   size=Pt(16.5), color=MAROON_DK)
    lp = doc.add_paragraph()
    lp.paragraph_format.first_line_indent = Inches(0)
    hline(lp)
    mula_buf = []
    last_body = [None]

    def flush_mula():
        if not mula_buf:
            return None
        pp = render_mula_group(doc, E, list(mula_buf))
        nums = re.findall(r'॥\s*([൦-൯]+)\s*॥', "\n".join(mula_buf))
        first_v = m2i(nums[0]) if nums else None
        mula_buf.clear()
        return pp, first_v

    pending_fig = [None]
    i = 0
    while i < len(md):
        line = md[i]
        i += 1
        s = line.strip()
        if not s:
            continue
        if s.startswith('```'):
            continue
        if s.startswith('# '):
            flush_mula()
            # file-internal H1 duplicates the chapter H1 on-page: skip
            # rendering entirely (no heading, no anchor). Scan mirrors this.
            continue
        if s.startswith('## '):
            flush_mula()
            t = s[3:].strip()
            if t.startswith('ഇതിൽനിന്ന്'):
                m = re.search(r'[൦-൯]+', t)
                bm = AG.recap(code, t)
                styled_heading(doc, t, 3, anchor=bm, size=Pt(11.0),
                               color=SLATE)
            else:
                m = re.search(r'[൦-൯]+', t)
                slug = re.sub(r'\s+', '_', re.sub(
                    r'[^\w൦-൯]+', ' ', t, flags=re.UNICODE)).strip('_')[:20]
                bm = AG.h2(code, t)
                styled_heading(doc, t, 2, anchor=bm,
                               size=Pt(12.5), color=MAROON)
            continue
        if s.startswith('### '):
            flush_mula()
            styled_heading(doc, s[4:].strip(), 3, size=Pt(11.0), color=SLATE)
            continue
        if s.startswith('>'):
            mula_buf.append(s)
            continue
        if mula_buf:
            res = flush_mula()
            if res:
                pp, first_v = res
                if pending_fig[0] is not None:
                    render_figure_unit(doc, E, pending_fig[0])
                    pending_fig[0] = None
                if first_v in CURATED_VERSES:
                    pending_fig[0] = first_v
        if s.startswith('|'):
            block = [s]
            while i < len(md) and md[i].strip().startswith('|'):
                block.append(md[i].strip())
                i += 1
            rows = [parse_md_row(t) for t in block]
            rows = [r for r in rows if not all(
                re.fullmatch(r':?---+:?', c or '---') for c in r)]
            if rows:
                add_word_table(doc, rows[0], rows[1:], Pt(E["body_pt"] - 1))
            continue
        if s.startswith('- ') or s.startswith('* '):
            pp = doc.add_paragraph()
            pp.paragraph_format.first_line_indent = Inches(0)
            pp.paragraph_format.left_indent = Inches(0.3)
            add_mixed_text(pp, "•  ", HEAD_FONT, None, True, False, GOLD)
            add_rich_mixed(pp, s[2:], BODY_FONT, None, INK)
            last_body[0] = pp
            continue
        if s.startswith('**കുറിപ്പ്'):
            # true footnote anchored at previous body paragraph (spec P)
            m = re.match(r'\*\*കുറിപ്പ്\s*([൦-൯]+)\s*—\*\*\s*(.*)', s)
            note_text = s
            if m:
                note_text = "**കുറിപ്പ് %s —** %s" % (m.group(1), m.group(2))
            if last_body[0] is not None:
                F.add(last_body[0], note_text)
            else:
                pp = body_para(doc, E, note_text, None, None,
                               no_indent=True, size=Pt(E["note_pt"]))
            continue
        is_lead = s.startswith(LEAD_LABELS[:3])
        pp = doc.add_paragraph()
        if is_lead:
            pp.paragraph_format.first_line_indent = Inches(0)
        if is_lead:
            m = re.match(r'(\*\*.+?:\*\*)\s*(.*)', s, re.S)
            if m:
                add_mixed_text(pp, m.group(1)[2:-3] + " — ", HEAD_FONT,
                               None, True, False, SLATE)
                add_rich_mixed(pp, m.group(2), BODY_FONT, None, INK)
            else:
                add_rich_mixed(pp, s, BODY_FONT, None, INK)
        else:
            add_rich_mixed(pp, s, BODY_FONT, None, INK)
        last_body[0] = pp
        if pending_fig[0] is not None:
            # figure unit directly after its verse-group commentary opens:
            # place after the FIRST body paragraph following the mula
            render_figure_unit(doc, E, pending_fig[0])
            pending_fig[0] = None
    if mula_buf:
        flush_mula()
    if pending_fig[0] is not None:
        render_figure_unit(doc, E, pending_fig[0])
        pending_fig[0] = None
    # keep closing ornament with its content (never strand "• • •" alone)
    if last_body[0] is not None:
        keep_with_next(last_body[0])
    ep = doc.add_paragraph()
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(ep, "• • •", HEAD_FONT, Pt(10), False, False, GOLD)


# ---------------- cover + frontmatter ----------------
def build_cover(doc, E):
    for _ in range(2):
        sp = doc.add_paragraph()
        sp.paragraph_format.first_line_indent = Inches(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(p, "തന്ത്രാലോകം · മലയാളപരമ്പര", HEAD_FONT, Pt(11),
                   False, False, GOLD)
    ep = doc.add_paragraph()
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep.paragraph_format.first_line_indent = Inches(0)
    ep.paragraph_format.space_before = Pt(2)
    shade(ep, "FAF4E6")
    box(ep, "A67C2E", "4")
    try:
        ep.add_run().add_picture(EMBLEM, width=E["emblem_w"])
    except Exception:
        pass
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.first_line_indent = Inches(0)
    t.paragraph_format.space_before = Pt(6)
    add_mixed_text(t, "തന്ത്രാലോകം", HEAD_FONT, Pt(30), True, False,
                   MAROON_DK)
    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t2.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(t2, "Tantrāloka · Malayalam Edition", HEAD_FONT, Pt(11),
                   False, True, SLATE)
    rl = doc.add_paragraph()
    rl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rl.paragraph_format.first_line_indent = Inches(0)
    hline(rl)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.first_line_indent = Inches(0)
    sub.paragraph_format.space_before = Pt(4)
    add_mixed_text(sub, "ഒന്നാം ഭാഗം · ഒന്നാം അദ്ധ്യായം", BODY_FONT,
                   Pt(13), True, False, SLATE)
    au = doc.add_paragraph()
    au.alignment = WD_ALIGN_PARAGRAPH.CENTER
    au.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(au, "അഭിനവഗുപ്തൻ — ജയരഥന്റെ വിവേകവ്യാഖ്യാനത്തോടെ",
                   BODY_FONT, Pt(11), False, False, INK)
    ed = doc.add_paragraph()
    ed.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ed.paragraph_format.first_line_indent = Inches(0)
    ed.paragraph_format.space_before = Pt(10)
    add_mixed_text(ed, "മൂലം മലയാളലിപിയിൽ · മലയാളവ്യാഖ്യാനം", HEAD_FONT,
                   Pt(9), False, False, GREY)
    fin = doc.add_paragraph()
    fin.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fin.paragraph_format.first_line_indent = Inches(0)
    fin.paragraph_format.space_before = Pt(14)
    add_mixed_text(fin, "• • •", HEAD_FONT, Pt(10), False, False, GOLD)
    imp = doc.add_paragraph()
    imp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imp.paragraph_format.first_line_indent = Inches(0)
    imp.paragraph_format.space_before = Pt(56)
    add_mixed_text(imp, "തന്ത്രാലോകം മലയാളപരമ്പര · ഒന്നാം ഭാഗം",
                   HEAD_FONT, Pt(8.5), False, False, GREY)


def render_intro_file(doc, E, fm):
    """ഗ്രന്ഥപരിചയം + chapter-overview + text-name index (ta_front.md)."""
    path = os.path.join(TR, "ta_front.md")
    with open(path, encoding="utf-8") as f:
        md = f.read().splitlines()
    seen_h1 = [False]
    for line in md:
        s = line.strip()
        if not s or s.startswith('# ') or s.startswith('```'):
            continue
        if s.startswith('## '):
            t = s[3:].strip()
            if t.startswith('എങ്ങനെ വായിക്കാം'):
                continue  # superseded by expanded guide
            if not seen_h1[0]:
                h = styled_heading(doc, "ഗ്രന്ഥപരിചയം", 1,
                                   anchor="fm_intro", size=Pt(16.5),
                                   color=MAROON_DK)
                fm.append((1, "ഗ്രന്ഥപരിചയം", "fm_intro"))
                seen_h1[0] = True
            styled_heading(doc, t, 2, anchor="fm_intro_%d" % len(fm),
                           size=Pt(12.5), color=MAROON)
            fm.append((2, t, "fm_intro_%d" % (len(fm))))
            continue
        if s.startswith('- ') or s.startswith('* '):
            pp = doc.add_paragraph()
            pp.paragraph_format.first_line_indent = Inches(0)
            pp.paragraph_format.left_indent = Inches(0.3)
            add_mixed_text(pp, "•  ", HEAD_FONT, None, True, False, GOLD)
            add_rich_mixed(pp, s[2:], BODY_FONT, None, INK)
            continue
        body_para(doc, E, s, no_indent=False)


def fm_heading(doc, text, fm, level=1, anchor=None):
    size = Pt(16.5) if level == 1 else Pt(12.5)
    color = MAROON_DK if level == 1 else MAROON
    h = styled_heading(doc, text, level, anchor=anchor, size=size,
                       color=color)
    if anchor:
        fm.append((level, text, anchor))
    return h


def build_frontmatter(doc, E, display, toc_pages=None):
    fm = []
    # half-title
    for _ in range(4):
        sp = doc.add_paragraph()
        sp.paragraph_format.first_line_indent = Inches(0)
    ht = doc.add_paragraph()
    ht.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ht.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(ht, "തന്ത്രാലോകം", HEAD_FONT, Pt(22), True, False,
                   MAROON_DK)
    ht2 = doc.add_paragraph()
    ht2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ht2.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(ht2, "ഒന്നാം ഭാഗം · ഒന്നാം അദ്ധ്യായം", BODY_FONT,
                   Pt(12), False, False, SLATE)
    doc.add_page_break()
    # title page
    for _ in range(2):
        sp = doc.add_paragraph()
        sp.paragraph_format.first_line_indent = Inches(0)
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(tp, "തന്ത്രാലോകം", HEAD_FONT, Pt(26), True, False,
                   MAROON_DK)
    tp2 = doc.add_paragraph()
    tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp2.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(tp2, "Tantrāloka — Volume One · Chapter One",
                   HEAD_FONT, Pt(10.5), False, True, SLATE)
    rl = doc.add_paragraph()
    rl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rl.paragraph_format.first_line_indent = Inches(0)
    hline(rl)
    au = doc.add_paragraph()
    au.alignment = WD_ALIGN_PARAGRAPH.CENTER
    au.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(au, "അഭിനവഗുപ്തൻ", BODY_FONT, Pt(14), True, False, INK)
    au2 = doc.add_paragraph()
    au2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    au2.paragraph_format.first_line_indent = Inches(0)
    add_mixed_text(au2, "ജയരഥന്റെ വിവേകം എന്ന വ്യാഖ്യാനത്തോടെ", BODY_FONT,
                   Pt(11), False, False, INK)
    ml = doc.add_paragraph()
    ml.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ml.paragraph_format.first_line_indent = Inches(0)
    ml.paragraph_format.space_before = Pt(8)
    add_mixed_text(ml, "മലയാളപതിപ്പ് · മൂലം മലയാളലിപിയിൽ", HEAD_FONT,
                   Pt(10), False, False, SLATE)
    doc.add_page_break()
    # edition note
    fm_heading(doc, "പതിപ്പുകുറിപ്പ്", fm, 1, "fm_edition")
    body_para(doc, E, "ഈ പതിപ്പ് മാർക്ക് ഡിക്സ്കോവ്സ്കിയുടെ ഇംഗ്ലീഷ് "
              "പരിഭാഷയെ (൨൦൨൩, ഒന്നാം ഭാഗം) ആധാരമാക്കിയുള്ള ചുരുക്കിയ "
              "മലയാളപതിപ്പാണ് — ഒന്നാം അദ്ധ്യായവും അനുബന്ധങ്ങളും. "
              "കാശ്മീരഗ്രന്ഥമാലയുടെ അച്ചടിമൂലം പിന്തുടരുന്നു.",
              no_indent=True)
    body_para(doc, E, "മൂലസംസ്കൃതം മൂലമായി തന്നെ — മലയാളലിപിയിൽ — "
              "സ്വർണ്ണപ്പെട്ടികളിൽ നിലനിർത്തിയിരിക്കുന്നു; മൂലപാഠം "
              "പരിഭാഷപ്പെടുത്തിയിട്ടില്ല. ഓരോ മൂലശ്ലോകത്തിനും തുടർന്ന് "
              "മലയാളത്തിൽ നേരർത്ഥവും വിശദീകരണവും നൽകിയിരിക്കുന്നു. "
              "വിവേകവ്യാഖ്യാനം ചുരുക്കമാണ് — സിദ്ധാന്തഭാഗം മാത്രം.",
              no_indent=True)
    body_para(doc, E, "സ്വകാര്യപഠനത്തിനുള്ള പരിഭാഷണം മാത്രം. "
              "ശ്ലോകസംഖ്യകൾ കാശ്മീരപാഠപ്രകാരം മലയാള അക്കങ്ങളിൽ.",
              no_indent=True)
    # reading guide
    fm_heading(doc, "എങ്ങനെ വായിക്കാം", fm, 1, "fm_reading")
    rg = [
        "മൂലം: സ്വർണ്ണപ്പെട്ടിയിലെ സംസ്കൃതം (മലയാളലിപിയിൽ) ചൊല്ലി, "
        "തുടർന്ന് നേരർത്ഥം വായിക്കുക.",
        "ആചാര്യവ്യാഖ്യാനം ജയരഥവിവേകത്തിന്റെ സിദ്ധാന്തച്ചുരുക്കമാണ്; "
        "വിശകലനം ആശയം തെളിയിക്കുന്നു.",
        "ചിത്രം അതാത് ഭാഗത്തിന്റെ ആശയഭേദം കാട്ടുന്നു; അടിക്കുറിപ്പ് "
        "ഏത് ശ്ലോകവുമായി ബന്ധപ്പെട്ടതെന്ന് പറയുന്നു.",
        "പ്രധാനപദങ്ങൾ ആദ്യപ്രയോഗത്തിൽ [ബ്രാക്കറ്റിൽ] വരുന്നു — "
        "ഉദാ. പരമശിവൻ [Paramaśiva].",
        "അടിക്കുറിപ്പുകൾ അതാത് ഭാഗത്തിന്റെ തുടർച്ചയായ "
        "കുറിപ്പുകളാണ്.",
    ]
    for item in rg:
        pp = doc.add_paragraph()
        pp.paragraph_format.first_line_indent = Inches(0)
        pp.paragraph_format.left_indent = Inches(0.3)
        add_mixed_text(pp, "•  ", HEAD_FONT, None, True, False, GOLD)
        add_rich_mixed(pp, item, BODY_FONT, None, INK)
    # book introduction from ta_front.md (its എങ്ങനെ-വായിക്കാം is
    # superseded by the expanded guide above — see CONTENT_CORRECTIONS)
    render_intro_file(doc, E, fm)
    # contents
    doc.add_page_break()
    fm_heading(doc, "ഉള്ളടക്കം", fm, 1, "fm_toc")
    for lvl, title, anchor in display:
        pp = doc.add_paragraph()
        pp.paragraph_format.first_line_indent = Inches(0)
        pp.paragraph_format.space_before = Pt(1) if lvl == 1 else Pt(0)
        pp.paragraph_format.space_after = Pt(1) if lvl == 1 else Pt(0)
        if lvl == 2:
            pp.paragraph_format.left_indent = Inches(0.3)
        add_toc_dot_leader(pp)
        size = Pt(10.5) if lvl == 1 else Pt(9.5)
        add_internal_link(pp, anchor, title, size)
        pg = (toc_pages or {}).get(anchor)
        if pg is not None:
            r = pp.add_run("\t%d" % pg)
            set_run_font(r, HEAD_FONT, size, False, False, GREY)
    return fm


# ---------------- appendices + colophon ----------------
def render_appendices(doc, E, F, AG):
    path = os.path.join(TR, "ta_app.md")
    with open(path, encoding="utf-8") as f:
        md = f.read().splitlines()
    styled_heading(doc, "അനുബന്ധങ്ങൾ", 1, anchor=AG.get("ch_APP"),
                   size=Pt(16.5), color=MAROON_DK)
    lp = doc.add_paragraph()
    lp.paragraph_format.first_line_indent = Inches(0)
    hline(lp)
    app1_paras = [0]
    in_app1 = [False]
    i = 0
    while i < len(md):
        line = md[i]
        i += 1
        s = line.strip()
        if not s or s.startswith('# ') or s.startswith('```'):
            continue
        if s.startswith('## '):
            t = s[3:].strip()
            if t.startswith('അനുബന്ധം'):
                parts = t.split('—', 1)
                kick = doc.add_paragraph()
                kick.alignment = WD_ALIGN_PARAGRAPH.LEFT
                kick.paragraph_format.first_line_indent = Inches(0)
                kick.paragraph_format.space_after = Pt(0)
                add_mixed_text(kick, parts[0].strip(), HEAD_FONT, Pt(9),
                               True, False, GOLD)
                code = 'APP1' if '൧' in parts[0] else (
                    'APP2' if '൨' in parts[0] else 'APPG')
                styled_heading(doc, parts[1].strip() if len(parts) > 1 else t,
                               1, anchor=AG.get(f"ch_{code}"),
                               size=Pt(16.5), color=MAROON_DK)
                in_app1[0] = (code == 'APP1')
                app1_paras[0] = 0
                rl = doc.add_paragraph()
                rl.paragraph_format.first_line_indent = Inches(0)
                hline(rl)
            else:
                styled_heading(doc, t, 2, anchor=AG.get("ch_APPG"),
                               size=Pt(12.5), color=MAROON)
            continue
        if s.startswith('|'):
            block = [s]
            while i < len(md) and md[i].strip().startswith('|'):
                block.append(md[i].strip())
                i += 1
            rows = [parse_md_row(t) for t in block]
            rows = [r for r in rows if not all(
                re.fullmatch(r':?---+:?', c or '---') for c in r)]
            if rows:
                if any('അളവ്' in c for c in rows[0]):
                    cap = "പട്ടിക ൧ — ദ്വാദശാന്തം: തലം · അധിപൻ · സ്ഥാനം · അളവ്"
                else:
                    cap = "പട്ടിക ൨ — ചുരുക്കപ്പേരുകൾ: ചുരുക്കം · വിപുലരൂപം · മലയാളം"
                add_word_table(doc, rows[0], rows[1:], Pt(E["body_pt"] - 1),
                               caption=cap)
            continue
        if re.match(r'^[൦-൯]+\.', s):
            pp = doc.add_paragraph()
            pp.paragraph_format.first_line_indent = Inches(0)
            add_rich_mixed(pp, s, BODY_FONT, None, INK)
            for r in pp.runs:
                r.bold = True
            app1_paras[0] += 0
            continue
        pp = body_para(doc, E, s, no_indent=False)
        if in_app1[0]:
            app1_paras[0] += 1
            if app1_paras[0] == 3:
                # appendix figure: twelve sound-levels ascent
                fp = os.path.join(FIGDIR, "fig_APP1.png")
                if os.path.exists(fp):
                    up = doc.add_paragraph()
                    up.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    up.paragraph_format.first_line_indent = Inches(0)
                    up.paragraph_format.space_before = Pt(4)
                    up.paragraph_format.space_after = Pt(1)
                    keep_lines(up)
                    try:
                        up.add_run().add_picture(fp, width=E["fig_w"])
                    except Exception:
                        pass
                    keep_with_next(up)
                    cp = doc.add_paragraph()
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cp.paragraph_format.first_line_indent = Inches(0)
                    cp.paragraph_format.space_after = Pt(6)
                    add_mixed_text(cp, "ചിത്രം — പന്ത്രണ്ട് ശബ്ദതലങ്ങളുടെ "
                                   "കയറ്റം · നാദതലം", HEAD_FONT,
                                   Pt(E["cap_pt"]), False, True, MAROON)
                    keep_lines(cp)
    # Appendix closes on its table (the table's bottom rule is the visual
    # close). A separate closing line + ornament orphaned onto a near-empty
    # page (verified A4 p122 / Trade p192); chapters keep their ornaments.


def build_colophon(doc, E):
    styled_heading(doc, "സമാപനം · കൊളോഫൺ", 1, anchor="ch_COL",
                   size=Pt(16.5), color=MAROON_DK)
    lp = doc.add_paragraph()
    lp.paragraph_format.first_line_indent = Inches(0)
    hline(lp)
    body_para(doc, E, "ഇതി അഭിനവഗുപ്തവിരചിതേ തന്ത്രാലോകനാമ്നി ഒന്നാം "
              "അദ്ധ്യായം മലയാളപതിപ്പ് സമ്പൂർണ്ണം — അനുബന്ധങ്ങൾ "
              "(ശബ്ദതലങ്ങൾ, ദ്വാദശാന്തം) സഹിതം. മൂലശ്ലോകങ്ങൾ സംസ്കൃതത്തിൽ "
              "മലയാളലിപിയിൽ നിലനിർത്തി; തുടർന്ന് മലയാളത്തിൽ നേരർത്ഥവും "
              "വിശദീകരണവും നൽകി; വിവേകവ്യാഖ്യാനം ചുരുക്കി സിദ്ധാന്തഭാഗം "
              "മാത്രം നൽകി.", no_indent=True)
    body_para(doc, E, "പരമ്പരശൈലി v02 — മൂലശ്ലോകങ്ങൾ ൩൩൪ (൧–൩൩൫, ൨൪൬ "
              "ഒഴികെ); ചിത്രങ്ങൾ ൩൬; അടിക്കുറിപ്പുകൾ ൧൩.",
              no_indent=True)
    fin = doc.add_paragraph()
    fin.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fin.paragraph_format.first_line_indent = Inches(0)
    fin.paragraph_format.space_before = Pt(12)
    add_mixed_text(fin, "ശുഭമസ്തു · ॥ ഓം തത്സത് ॥", BODY_FONT, Pt(11),
                   True, False, MAROON)


# ---------------- driver ----------------
def build(edition, toc_pages=None, starts=None):
    E = EDITIONS[edition]
    doc = Document()
    style_document(doc, E)
    # NOTE: pages stay white (w:background page color was tried and is
    # ignored by the LibreOffice converter, which would desync DOCX vs PDF;
    # warmth comes from the ivory plates: emblem, mula boxes, figures).
    F = Footnotes(doc, Pt(E["note_pt"]))
    AG = AnchorGen()
    display, _all = scan_headings(AnchorGen())  # TOC first (spec M)
    # cover (no folio; framed emblem plate, ivory ground)
    sec0 = doc.sections[0]
    setup_section(sec0, E, folio=False)
    build_cover(doc, E)
    # frontmatter (roman)
    doc.add_section(WD_SECTION.NEW_PAGE)
    sec_fm = doc.sections[-1]
    setup_section(sec_fm, E, footer_fmt="lowerRoman", footer_start=1)
    fm = build_frontmatter(doc, E, display, toc_pages)
    # chapters
    for idx, (fname, title, short, code) in enumerate(ORDER):
        if code == "APP":
            continue
        doc.add_section(WD_SECTION.NEW_PAGE)
        sec = doc.sections[-1]
        dflt = 1 if idx == 0 else None
        start = (starts or {}).get(code, dflt)
        setup_section(sec, E,
                      header_odd=E["series_head"] + " · " + short,
                      header_even=None,
                      footer_fmt="decimal", footer_start=start,
                      restart_footnotes=True)
        render_chapter_file(doc, E, fname, title, short, code, F, AG)
    # appendices
    doc.add_section(WD_SECTION.NEW_PAGE)
    sec_ap = doc.sections[-1]
    setup_section(sec_ap, E,
                  header_odd=E["series_head"] + " · അനുബന്ധങ്ങൾ",
                  header_even=None,
                  footer_fmt="decimal",
                  footer_start=(starts or {}).get("APP"),
                  restart_footnotes=True)
    render_appendices(doc, E, F, AG)
    # colophon
    doc.add_section(WD_SECTION.NEW_PAGE)
    sec_co = doc.sections[-1]
    setup_section(sec_co, E,
                  header_odd=E["series_head"] + " · സമാപനം",
                  header_even=None,
                  footer_fmt="decimal",
                  footer_start=(starts or {}).get("COL"))
    build_colophon(doc, E)
    F.finalize()
    # build gate: renderer anchors must equal pre-scanned TOC anchors
    scanned = set(_all)
    if set(AG.used) != scanned:
        missing = sorted(scanned - set(AG.used))
        extra = sorted(set(AG.used) - scanned)
        raise SystemExit("ANCHOR MISMATCH missing=%s extra=%s"
                         % (missing, extra))
    toc = fm + display
    doc.core_properties.title = "തന്ത്രാലോകം — ഒന്നാം ഭാഗം · ഒന്നാം അദ്ധ്യായം (മലയാളപതിപ്പ് v02)"
    doc.core_properties.author = "അഭിനവഗുപ്തൻ (മൂലം); ജയരഥൻ (വിവേകം); മലയാളപതിപ്പ്"
    doc.core_properties.subject = "തന്ത്രാലോകം ഒന്നാം അദ്ധ്യായം — മലയാളപതിപ്പ്, പരമ്പര ശൈലി v02"
    doc.core_properties.keywords = "തന്ത്രാലോകം, അഭിനവഗുപ്തൻ, ജയരഥൻ, കാശ്മീർ ശൈവദർശനം, മലയാളം"
    doc.core_properties.language = "ml"
    doc.core_properties.comments = ("Tantraloka Malayalam Series v02 house style; "
                                    "mula in Malayalam script with Malayalam commentary")
    out = os.path.join(BASE, "pdf",
                       "Tantraloka_Malayalam_Volume1_v02_%s.docx" % E["suffix"])
    doc.save(out)
    print("saved %s (%d KB)" % (out, os.path.getsize(out) // 1024))
    return out, toc


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--edition", choices=["a4", "trade"], default="a4")
    ap.add_argument("--toc", default=None,
                    help="JSON anchor->page for pass 2")
    ap.add_argument("--starts", default=None,
                    help="JSON section-code->folio-start for pass 2")
    args = ap.parse_args()
    toc_pages = None
    if args.toc:
        with open(args.toc, encoding="utf-8") as f:
            toc_pages = json.load(f)
    starts = None
    if args.starts:
        with open(args.starts, encoding="utf-8") as f:
            starts = json.load(f)
    out, toc = build(args.edition, toc_pages, starts)
    base = os.path.splitext(os.path.basename(out))[0]
    with open(os.path.join(BASE, "pdf", base + ".anchors.json"), "w",
              encoding="utf-8") as f:
        json.dump([{"level": l, "title": t, "anchor": a}
                   for l, t, a in toc], f, ensure_ascii=False, indent=1)
    print("anchors:", len(toc))


if __name__ == "__main__":
    main()
